"""Excel: экспорт заказов/клиентов, импорт товаров, парсер накладных."""
import logging
import re
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="2D6A4F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


# ═══════════════════════════════════════════════════════════
# ОБЩИЕ УТИЛИТЫ
# ═══════════════════════════════════════════════════════════

def _header(ws) -> None:
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN


def _autosize(ws, max_width: int = 40) -> None:
    for col in ws.columns:
        ml = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(ml + 3, max_width)


def _stream(wb: Workbook) -> BytesIO:
    s = BytesIO()
    wb.save(s)
    s.seek(0)
    return s


# ═══════════════════════════════════════════════════════════
# ЭКСПОРТ ЗАКАЗОВ И КЛИЕНТОВ
# ═══════════════════════════════════════════════════════════

def export_orders_to_excel(orders: list) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    ws.append(["№", "Дата", "Компания", "Контакт", "Email",
               "Товар", "Упак.", "Шт/упак", "Всего шт",
               "Цена/шт", "Сумма позиции",
               "Подытог", "Скидка %", "Итого", "Статус", "Комментарий"])
    _header(ws)
    for order in orders:
        user = order.user
        first = True
        if not order.items:
            ws.append([order.id,
                       order.created_at.strftime("%d.%m.%Y %H:%M"),
                       user.company_name if user else "Гость",
                       user.full_name if user else "",
                       user.email if user else "",
                       "—", 0, 0, 0, 0, 0,
                       order.subtotal, order.discount_percent,
                       order.total_price, order.status, order.comment])
            continue
        for item in order.items:
            pack = item.package_size or 1
            stems = item.quantity * pack
            ws.append([
                order.id if first else "",
                order.created_at.strftime("%d.%m.%Y %H:%M") if first else "",
                (user.company_name if user else "Гость") if first else "",
                (user.full_name if user else "") if first else "",
                (user.email if user else "") if first else "",
                item.product_name,
                item.quantity,
                pack,
                stems,
                item.price,
                item.subtotal,
                order.subtotal if first else "",
                order.discount_percent if first else "",
                order.total_price if first else "",
                order.status if first else "",
                order.comment if first else "",
            ])
            first = False
    _autosize(ws)
    return _stream(wb)


def export_customers_to_excel(customers: list[dict]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Клиенты"
    ws.append(["Компания", "Контакт", "Email", "Телефон", "Скидка %",
               "Заказов", "Сумма, ₽", "Последний заказ"])
    _header(ws)
    for c in customers:
        ws.append([c.get("company_name", ""), c.get("full_name", ""),
                   c.get("email", ""), c.get("phone", ""),
                   c.get("discount_percent", 0),
                   c.get("orders_count", 0), c.get("total_sum", 0),
                   c.get("last_order_date", "")])
    _autosize(ws)
    return _stream(wb)


# ═══════════════════════════════════════════════════════════
# ИМПОРТ СПРАВОЧНИКА ТОВАРОВ
# ═══════════════════════════════════════════════════════════

PRODUCT_IMPORT_HEADERS = ["Название", "Страна", "Длина, см",
                          "Единица", "В упаковке", "Мин. заказ",
                          "Категория", "Описание"]


def import_products_from_excel(content: bytes) -> tuple[list[dict], list[str]]:
    warnings: list[str] = []
    products: list[dict] = []
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as e:
        return [], [f"Не удалось открыть файл: {e}"]

    ws = wb.active
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(c is None for c in row):
            continue
        try:
            name = str(row[0]).strip() if row[0] else ""
            if not name:
                warnings.append(f"Строка {i}: пустое название")
                continue

            def safe(idx, default=""):
                return row[idx] if idx < len(row) else None

            products.append({
                "name": name,
                "country": str(safe(1) or "").strip(),
                "length_cm": int(safe(2) or 0),
                "unit": str(safe(3) or "упаковка").strip(),
                "package_size": int(safe(4) or 1),
                "min_quantity": int(safe(5) or 1),
                "category": str(safe(6) or "Прочее").strip(),
                "description": str(safe(7) or "").strip(),
                "image_url": "",
            })
        except (ValueError, TypeError) as e:
            warnings.append(f"Строка {i}: {e}")
    return products, warnings


def build_products_import_template() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(PRODUCT_IMPORT_HEADERS)
    ws.append(["Роза Freedom 60 см", "Эквадор", 60, "упаковка",
               25, 1, "Роза Эквадор", "Классическая красная"])
    _header(ws)
    _autosize(ws, max_width=30)
    return _stream(wb)


# ═══════════════════════════════════════════════════════════
# ПАРСЕР НАКЛАДНЫХ
# ═══════════════════════════════════════════════════════════

_COL_ALIASES: dict[str, list[str]] = {
    "name":         ["товар", "наименование", "название"],
    "quantity":     ["количество", "кол-во", "колво", "кол."],
    "unit":         ["ед.", "единица", "ед. изм", "ед.изм"],
    "price":        ["цена"],
    "total":        ["сумма", "стоимость"],
    "country":      ["страна происхождения", "страна"],
    "package_size": ["в упаковке", "упаковка", "шт в упак", "шт/упак"],
}

_STOP_WORDS = [
    "итого", "всего наименований", "в том числе",
    "ндс:", "подпись", "отпустил", "получил",
]

_LENGTH_MIN = 20
_LENGTH_MAX = 150

_CATEGORY_PATTERNS: list[tuple[str, str]] = [
    (r"гвоздик", "Гвоздика"),
    (r"хризантем", "Хризантема"),
    (r"тюльпан", "Тюльпан"),
    (r"пион", "Пион"),
    (r"горшеч", "Горшечные"),
    (r"зелень", "Зелень"),
    (r"\btessa\b", "Роза Эквадор"),
    (r"роза", "Роза Эквадор"),
]

# Предполагаемый размер упаковки по названию (если в накладной нет колонки)
_DEFAULT_PACK_BY_NAME: list[tuple[str, int]] = [
    (r"гвоздик", 20),
    (r"тюльпан", 50),
    (r"пион", 10),
    (r"хризантем", 10),
    (r"\btessa\b", 25),
    (r"роза", 25),
]


def guess_package_size(name: str) -> int:
    """Пытается угадать размер упаковки по названию товара."""
    low = name.lower()
    for pattern, size in _DEFAULT_PACK_BY_NAME:
        if re.search(pattern, low):
            return size
    return 1


def _read_xls(content: bytes) -> list[list]:
    import xlrd
    wb = xlrd.open_workbook(file_contents=content)
    sheet = wb.sheet_by_index(0)
    return [sheet.row_values(r) for r in range(sheet.nrows)]


def _read_xlsx(content: bytes) -> list[list]:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _to_float(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("\xa0", "")
    m = re.search(r"-?\d+(?:[.,]\d+)?", s)
    if not m:
        return 0.0
    try:
        return float(m.group().replace(",", "."))
    except ValueError:
        return 0.0


def _to_int(v) -> int:
    return int(_to_float(v))


def _clean_name(raw) -> str:
    if raw is None:
        return ""
    return re.sub(r"\s+", " ", str(raw)).strip()


def _find_header(rows: list[list]) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(rows[:80]):
        col_map: dict[str, int] = {}
        for j, cell in enumerate(row):
            s = _clean_name(cell).lower().replace("\n", " ")
            if not s:
                continue
            for field, aliases in _COL_ALIASES.items():
                if field in col_map:
                    continue
                for alias in aliases:
                    if alias in s:
                        col_map[field] = j
                        break
        if {"name", "quantity", "price"} <= col_map.keys():
            return i, col_map
    return -1, {}


def _is_stop_row(row: list) -> bool:
    if not row:
        return False
    chunk = " ".join(_clean_name(c) for c in row).lower().strip()
    if not chunk:
        return False
    return any(w in chunk for w in _STOP_WORDS)


def _extract_length(name: str) -> int:
    for m in re.finditer(r"\b(\d{2,3})\b", name):
        n = int(m.group(1))
        if _LENGTH_MIN <= n <= _LENGTH_MAX:
            return n
    return 0


def _guess_category(name: str) -> str:
    low = name.lower()
    for pattern, category in _CATEGORY_PATTERNS:
        if re.search(pattern, low):
            return category
    return "Прочее"


def parse_invoice(content: bytes, filename: str
                  ) -> tuple[list[dict], list[str]]:
    """
    Универсальный парсер накладных.

    Возвращает:
        items — [{"name", "quantity", "unit", "price", "total",
                  "length_cm", "category", "country", "package_size"}]
        warnings
    """
    warnings: list[str] = []
    fn = (filename or "").lower()

    try:
        if fn.endswith(".xls"):
            rows = _read_xls(content)
        elif fn.endswith(".xlsx"):
            rows = _read_xlsx(content)
        else:
            return [], ["Нужен файл .xls или .xlsx"]
    except Exception as e:
        return [], [f"Не удалось открыть файл: {e}"]

    if not rows:
        return [], ["Файл пуст"]

    header_idx, col_map = _find_header(rows)
    if header_idx < 0:
        return [], [
            "Не нашёл шапку таблицы. Нужны колонки: "
            "«Товар», «Количество», «Цена»"
        ]

    logger.info("Накладная: шапка в строке %d, колонки %s",
                header_idx + 1, col_map)

    items: list[dict] = []

    def cell(row, field):
        j = col_map.get(field)
        if j is None or j >= len(row):
            return None
        return row[j]

    for i in range(header_idx + 1, len(rows)):
        row = rows[i]

        if _is_stop_row(row):
            logger.info("Накладная: стоп на строке %d", i + 1)
            break

        name = _clean_name(cell(row, "name"))
        if len(name) < 2:
            continue

        if any(w in name.lower() for w in ["итого", "всего", "ндс"]):
            continue

        quantity = _to_int(cell(row, "quantity"))
        price = _to_float(cell(row, "price"))
        total = _to_float(cell(row, "total"))

        if quantity <= 0 or price <= 0:
            warnings.append(
                f"Строка {i + 1}: пропущена "
                f"(кол-во={quantity}, цена={price})"
            )
            continue

        calc = quantity * price
        if total and abs(total - calc) > 1:
            warnings.append(
                f"Строка {i + 1}: сумма {total} ≠ "
                f"{quantity}×{price}={calc:.2f}"
            )

        unit = _clean_name(cell(row, "unit")) or "шт"
        country = _clean_name(cell(row, "country"))
        pack_from_invoice = _to_int(cell(row, "package_size")) or 0

        items.append({
            "name": name,
            "quantity": quantity,           # ШТУК
            "unit": unit,
            "price": price,                 # за ШТУКУ
            "total": total,
            "length_cm": _extract_length(name),
            "category": _guess_category(name),
            "country": country,
            "package_size": pack_from_invoice,  # 0 = неизвестно
        })

    if not items:
        warnings.append("В таблице не нашлось строк с товарами")
        logger.warning(
            "Накладная: 0 позиций. header_idx=%d, col_map=%s",
            header_idx, col_map,
        )

    return items, warnings