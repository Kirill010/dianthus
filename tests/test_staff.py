# tests/test_staff.py
"""Сотрудники: роль «менеджер», главная админки, карточка клиента, листы сборки,
сброс пароля, настройки магазина, письмо «забыли пароль», счёт PDF."""
import re
from datetime import datetime, timedelta
from io import BytesIO

from app.config import config
from app.models import Claim, Order, PasswordReset, User
from app.services.site_settings import get_setting
from app.timeutils import shop_tz
from tests.helpers import csrf_from, login, make_order, make_user, seed_item


def as_manager(web, mem_db):
    make_user(mem_db, email="m@test.ru", password="Manager123", is_manager=True)
    login(web, "m@test.ru", "Manager123")


def as_admin(web, mem_db):
    make_user(mem_db, email="boss@test.ru", password="Admin1234", admin=True)
    login(web, "boss@test.ru", "Admin1234")


def today():
    return datetime.now(shop_tz()).date()


# ───────────────────────── права менеджера ─────────────────────────

def test_manager_sees_orders_but_not_money_or_admin_tabs(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru")
    make_order(mem_db, buyer, seed_item(mem_db, stock=20), packs=2)
    as_manager(web, mem_db)
    page = web.get("/admin")
    assert page.status_code == 200
    text = page.text
    assert "#1" in text and "(менеджер)" in text
    assert "Выручка" not in text                                  # финансы — только владельцу
    for hidden in ('href="#products"', 'href="#supplies"', 'href="#active"',
                   "/admin/integration", "/admin/legal", "/admin/settings", "/admin/audit",
                   "/admin/customers/export", "importModal"):
        assert hidden not in text, hidden
    assert "/admin/picking" in text and "/admin/claims" in text


def test_manager_can_work_with_orders_and_clients(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru", approved=False)
    order = make_order(mem_db, buyer, seed_item(mem_db, stock=20), packs=1)
    as_manager(web, mem_db)
    token = csrf_from(web.get("/admin").text)
    r = web.post("/admin/update_order_status",
                 data={"csrf_token": token, "order_id": order.id, "status": "Подтверждён"},
                 follow_redirects=False)
    assert r.status_code == 303
    mem_db.expire_all()
    assert mem_db.get(Order, order.id).status == "Подтверждён"
    web.post(f"/admin/orders/{order.id}/note", data={"csrf_token": token, "staff_note": "Позвонить"})
    mem_db.expire_all()
    assert mem_db.get(Order, order.id).staff_note == "Позвонить"
    web.post(f"/admin/users/{buyer.id}/approve", data={"csrf_token": token})
    mem_db.expire_all()
    assert mem_db.get(User, buyer.id).is_approved is True
    for url in ("/admin/orders/export", f"/admin/orders/{order.id}", f"/admin/orders/{order.id}/print",
                f"/admin/clients/{buyer.id}", "/admin/picking", f"/admin/orders/{order.id}/invoice"):
        assert web.get(url).status_code == 200, url


def test_manager_is_kept_out_of_everything_else(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru")
    as_manager(web, mem_db)
    token = csrf_from(web.get("/admin").text)
    for url in ("/admin/integration", "/admin/legal", "/admin/settings", "/admin/audit",
                "/admin/customers/export", "/admin/products/new", "/admin/supplies/new",
                "/admin/data-requests", "/admin/products/import/template"):
        r = web.get(url, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/catalog", url
    for url, data in (("/admin/users/%d/discount" % buyer.id, {"discount_percent": "50"}),
                      ("/admin/users/%d/make_admin" % buyer.id, {}),
                      ("/admin/users/%d/reject" % buyer.id, {}),
                      ("/admin/users/%d/manager" % buyer.id, {}),
                      ("/admin/users/%d/reset_password" % buyer.id, {}),
                      ("/admin/settings", {"payment_terms": "x", "preorder_lead_days": "a=1"})):
        r = web.post(url, data={"csrf_token": token, **data}, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/catalog", url
    mem_db.expire_all()
    u = mem_db.get(User, buyer.id)
    assert u.discount_percent == 0 and not u.is_admin and u.anonymized_at is None


def test_regular_client_has_no_staff_access_at_all(web, mem_db):
    make_user(mem_db, "c@test.ru")
    login(web, "c@test.ru", "Client123")
    for url in ("/admin", "/admin/picking", "/admin/clients/1", "/admin/orders/1/print"):
        r = web.get(url, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/catalog", url


def test_admin_grants_and_takes_manager_role(web, mem_db):
    helper = make_user(mem_db, "h@test.ru")
    pending = make_user(mem_db, "p@test.ru", approved=False)
    as_admin(web, mem_db)
    token = csrf_from(web.get(f"/admin/clients/{helper.id}").text)
    web.post(f"/admin/users/{helper.id}/manager", data={"csrf_token": token})
    web.post(f"/admin/users/{pending.id}/manager", data={"csrf_token": token})
    mem_db.expire_all()
    assert mem_db.get(User, helper.id).is_manager is True
    assert mem_db.get(User, pending.id).is_manager is False          # неодобренному роль не выдаём
    web.post(f"/admin/users/{helper.id}/manager", data={"csrf_token": token})
    mem_db.expire_all()
    assert mem_db.get(User, helper.id).is_manager is False


# ───────────────────────── главная админки ─────────────────────────

def test_admin_dashboard_todo_counts_and_urgent_filter(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru")
    make_user(mem_db, "pending@test.ru", approved=False)
    item = seed_item(mem_db, stock=100)
    soon = make_order(mem_db, buyer, item, packs=1)
    soon.desired_date = datetime(today().year, today().month, today().day) + timedelta(days=1)
    late = make_order(mem_db, buyer, item, packs=1)
    late.desired_date = datetime(today().year, today().month, today().day) - timedelta(days=2)
    far = make_order(mem_db, buyer, item, packs=1)
    far.desired_date = datetime(today().year, today().month, today().day) + timedelta(days=20)
    done = make_order(mem_db, buyer, item, packs=1, status="Выполнен")
    done.desired_date = datetime(today().year, today().month, today().day)
    mem_db.add(Claim(user_id=buyer.id, product_name="Роза", description="плохо плохо плохо"))
    mem_db.commit()
    as_admin(web, mem_db)
    page = web.get("/admin")
    assert "Что сделать сегодня" in page.text and "просрочен" in page.text
    ids = lambda html: set(re.findall(r"/admin/orders/(\d+)\"", html))
    urgent = web.get("/admin?due=1")
    assert ids(urgent.text) == {str(soon.id), str(late.id)}        # без «далёкого» и без выполненного
    assert "Выручка" in page.text


def test_status_chips_show_counts(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru")
    item = seed_item(mem_db, stock=100)
    make_order(mem_db, buyer, item, status="Новый")
    make_order(mem_db, buyer, item, status="Новый")
    make_order(mem_db, buyer, item, status="Отправлен")
    as_admin(web, mem_db)
    page = web.get("/admin").text
    assert "Новый (2)" in page and "Отправлен (1)" in page and "Выполнен (0)" in page


def test_dashboard_search_by_phone(web, mem_db):
    make_order(mem_db, make_user(mem_db, "a@test.ru", phone="+79990001122"), seed_item(mem_db, stock=9))
    as_admin(web, mem_db)
    assert "#1" in web.get("/admin", params={"q": "9990001122"}).text


# ───────────────────────── карточка клиента ─────────────────────────

def test_client_card_numbers_and_note(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru", company_name="ООО Ромашка")
    item = seed_item(mem_db, stock=100, price=40, pack=25)          # заказ на 1000 ₽
    make_order(mem_db, buyer, item, packs=1)
    make_order(mem_db, buyer, item, packs=2)
    make_order(mem_db, buyer, item, packs=5, status="Отменён")      # не считается
    as_admin(web, mem_db)
    page = web.get(f"/admin/clients/{buyer.id}")
    assert page.status_code == 200 and "ООО Ромашка" in page.text
    assert "3000 ₽" in page.text and "1500 ₽" in page.text           # сумма и средний чек без отменённого
    assert "нет отметки" in page.text                                # согласие на данные не давал
    token = csrf_from(page.text)
    web.post(f"/admin/clients/{buyer.id}/note", data={"csrf_token": token, "staff_note": "Любит розы 60"})
    mem_db.expire_all()
    assert mem_db.get(User, buyer.id).staff_note == "Любит розы 60"
    assert "Любит розы 60" in web.get(f"/admin/clients/{buyer.id}").text
    assert web.get("/admin/clients/9999").status_code == 404


def test_staff_note_is_never_shown_to_the_client(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru", staff_note="ТОЛЬКО ДЛЯ СОТРУДНИКОВ")
    order = make_order(mem_db, buyer, seed_item(mem_db, stock=9))
    order.staff_note = "ВНУТРЕННЕЕ О ЗАКАЗЕ"
    mem_db.commit()
    login(web, "b@test.ru", "Client123")
    for url in ("/profile", "/orders", "/catalog", "/cart", f"/orders/{order.id}/invoice"):
        r = web.get(url)
        assert "ТОЛЬКО ДЛЯ СОТРУДНИКОВ" not in r.text and "ВНУТРЕННЕЕ О ЗАКАЗЕ" not in r.text, url
    assert "ТОЛЬКО" not in web.get("/profile/data.json").text


# ───────────────────────── сброс пароля ─────────────────────────

def test_admin_resets_password_and_client_can_log_in_with_new_one(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru")
    as_admin(web, mem_db)
    token = csrf_from(web.get(f"/admin/clients/{buyer.id}").text)
    r = web.post(f"/admin/users/{buyer.id}/reset_password", data={"csrf_token": token})
    assert r.status_code == 200 and "no-store" in r.headers["cache-control"]
    temp = re.search(r'id="tmp-pwd"[^>]*value="([^"]+)"', r.text).group(1)
    assert len(temp) == 10 and any(c.isdigit() for c in temp) and any(c.isalpha() for c in temp)
    assert temp not in web.get(f"/admin/clients/{buyer.id}").text      # повторно не показывается
    web.cookies.clear()
    assert login(web, "b@test.ru", "Client123").headers["location"] == "/login"
    web.cookies.clear()
    assert login(web, "b@test.ru", temp).headers["location"] == "/catalog"


def test_reset_password_is_logged_without_the_password(web, mem_db):
    from app.models import AuditLog
    buyer = make_user(mem_db, "b@test.ru")
    as_admin(web, mem_db)
    token = csrf_from(web.get(f"/admin/clients/{buyer.id}").text)
    r = web.post(f"/admin/users/{buyer.id}/reset_password", data={"csrf_token": token})
    temp = re.search(r'id="tmp-pwd"[^>]*value="([^"]+)"', r.text).group(1)
    details = " ".join(a.details for a in mem_db.query(AuditLog).all())
    assert "Сброшен пароль" in details and temp not in details


# ───────────────────────── «забыли пароль» ─────────────────────────

def smtp_on(monkeypatch):
    for attr, val in (("SMTP_HOST", "smtp.test"), ("SMTP_USER", "u"), ("SMTP_PASSWORD", "p"),
                      ("SMTP_TO", "owner@test.ru"), ("SMTP_FROM", "shop@test.ru")):
        monkeypatch.setattr(config, attr, val)
    sent = []
    from app.services import notifier
    monkeypatch.setattr(notifier, "_smtp_send_raw", lambda cfg, msg, rcpt: sent.append((rcpt, msg.as_string())))
    from app.routers import password_reset
    monkeypatch.setattr(password_reset, "send_email_to",
                        lambda to, subject, html, text=None: sent.append((to, html)))
    return sent


def ask_reset(web, email):
    token = csrf_from(web.get("/forgot").text)
    return web.post("/forgot", data={"csrf_token": token, "email": email}, follow_redirects=True)


def test_forgot_password_full_flow(web, mem_db, monkeypatch):
    sent = smtp_on(monkeypatch)
    make_user(mem_db, "b@test.ru")
    r = ask_reset(web, "b@test.ru")
    assert "Если такой адрес зарегистрирован" in r.text
    assert len(sent) == 1 and sent[0][0] == "b@test.ru"
    link = re.search(r"/reset/([A-Za-z0-9_\-]+)", sent[0][1]).group(1)
    row = mem_db.query(PasswordReset).one()
    assert row.token_hash != link and len(row.token_hash) == 64        # токен в базе — только хэш

    page = web.get(f"/reset/{link}")
    assert page.status_code == 200 and 'name="password2"' in page.text
    token = csrf_from(page.text)
    bad = web.post(f"/reset/{link}", data={"csrf_token": token, "password": "abc", "password2": "abc"},
                   follow_redirects=True)
    assert "минимум 8" in bad.text.lower() or "Пароль" in bad.text
    ok = web.post(f"/reset/{link}", data={"csrf_token": token, "password": "NewPass123", "password2": "NewPass123"},
                  follow_redirects=False)
    assert ok.status_code == 303 and ok.headers["location"] == "/login"
    assert login(web, "b@test.ru", "NewPass123").headers["location"] == "/catalog"
    # одноразовая
    web.cookies.clear()
    again = web.get(f"/reset/{link}")
    assert "недействительна" in again.text


def test_forgot_password_does_not_reveal_unknown_emails(web, mem_db, monkeypatch):
    sent = smtp_on(monkeypatch)
    known = ask_reset(web, "nobody@test.ru").text
    make_user(mem_db, "b@test.ru")
    exists = ask_reset(web, "b@test.ru").text
    strip = lambda t: re.sub(r'name="csrf_token" value="[^"]+"|content="[^"]+"|nonce="[^"]+"', "", t)
    assert strip(known) == strip(exists)                             # одинаковый ответ
    assert len(sent) == 1 and mem_db.query(PasswordReset).count() == 1


def test_forgot_password_expired_and_replaced_links(web, mem_db, monkeypatch):
    sent = smtp_on(monkeypatch)
    make_user(mem_db, "b@test.ru")
    ask_reset(web, "b@test.ru")
    first = re.search(r"/reset/([A-Za-z0-9_\-]+)", sent[0][1]).group(1)
    ask_reset(web, "b@test.ru")
    second = re.search(r"/reset/([A-Za-z0-9_\-]+)", sent[1][1]).group(1)
    assert "недействительна" in web.get(f"/reset/{first}").text       # старая погашена
    assert 'name="password2"' in web.get(f"/reset/{second}").text
    row = mem_db.query(PasswordReset).filter(PasswordReset.used_at.is_(None)).one()
    row.expires_at = datetime.utcnow() - timedelta(minutes=1)
    mem_db.commit()
    assert "недействительна" in web.get(f"/reset/{second}").text


def test_forgot_password_without_mail_tells_to_call(web, mem_db, monkeypatch):
    for attr in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "SMTP_TO"):
        monkeypatch.setattr(config, attr, "")
    page = web.get("/forgot").text
    assert "Позвоните менеджеру" in page and 'name="email"' not in page
    assert mem_db.query(PasswordReset).count() == 0


def test_forgot_password_is_rate_limited(web, mem_db, monkeypatch):
    smtp_on(monkeypatch)
    make_user(mem_db, "b@test.ru")
    token = csrf_from(web.get("/forgot").text)
    codes = [web.post("/forgot", data={"csrf_token": token, "email": "b@test.ru"},
                      follow_redirects=False).status_code for _ in range(7)]
    assert 429 in codes and mem_db.query(PasswordReset).count() <= 3


def test_login_page_links_to_forgot(web):
    assert 'href="/forgot"' in web.get("/login").text


# ───────────────────────── листы сборки ─────────────────────────

def test_picking_list_sums_products_and_filters_by_date(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru", company_name="ООО Ромашка")
    rose = seed_item(mem_db, name="Роза Фридом", stock=100, pack=25)
    tulip = seed_item(mem_db, name="Тюльпан", stock=100, pack=50)
    d1 = datetime(today().year, today().month, today().day) + timedelta(days=3)
    a = make_order(mem_db, buyer, rose, packs=2)
    b = make_order(mem_db, buyer, rose, packs=3, status="В работе")
    c = make_order(mem_db, buyer, tulip, packs=1)
    make_order(mem_db, buyer, rose, packs=9, status="Выполнен")      # в лист не попадает
    a.desired_date = b.desired_date = d1
    mem_db.commit()
    as_admin(web, mem_db)
    page = web.get("/admin/picking").text
    assert "Роза Фридом" in page and ">5<" in page                   # 2 + 3 упаковки
    assert "Тюльпан" in page and "Выполнен" not in page.split("По заказам")[0]
    dated = web.get("/admin/picking", params={"date": d1.date().isoformat()}).text
    assert "Роза Фридом" in dated and "Тюльпан" not in dated
    assert f"/admin/orders/{c.id}" not in dated


def test_order_print_sheet(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru", company_name="ООО Ромашка")
    order = make_order(mem_db, buyer, seed_item(mem_db, name="Роза Фридом", stock=50, pack=25), packs=2)
    order.delivery_method, order.delivery_address = "delivery", "Москва, Цветочная 5"
    order.comment, order.staff_note = "Позвонить", "Достать свежие"
    mem_db.commit()
    as_admin(web, mem_db)
    page = web.get(f"/admin/orders/{order.id}/print").text
    for needle in ("Лист сборки", "Роза Фридом", "Москва, Цветочная 5", "Позвонить", "Достать свежие", "ООО Ромашка"):
        assert needle in page, needle
    assert web.get("/admin/orders/9999/print").status_code == 404


# ───────────────────────── настройки и счёт ─────────────────────────

def test_settings_are_saved_and_used_everywhere(web, mem_db):
    buyer = make_user(mem_db, "b@test.ru")
    as_admin(web, mem_db)
    page = web.get("/admin/settings")
    assert page.status_code == 200 and "Эквадор=10" in page.text
    token = csrf_from(page.text)
    r = web.post("/admin/settings", data={"csrf_token": token, "payment_terms": "Оплата в течение 3 дней",
                                          "preorder_lead_days": "Эквадор=14; Колумбия=9"}, follow_redirects=False)
    assert r.status_code == 303
    assert get_setting(mem_db, "payment_terms") == "Оплата в течение 3 дней"
    assert "за 14 дн." in web.get("/terms").text and "Оплата в течение 3 дней" in web.get("/terms").text
    assert "Эквадор=14, Колумбия=9" in web.get("/admin/settings").text


def test_settings_validation(web, mem_db):
    as_admin(web, mem_db)
    token = csrf_from(web.get("/admin/settings").text)
    for terms, days in (("", "Эквадор=10"), ("ок", "ерунда"), ("ок", "Эквадор=10, мусор")):
        web.post("/admin/settings", data={"csrf_token": token, "payment_terms": terms, "preorder_lead_days": days})
        assert "Предоплата 100%" in get_setting(mem_db, "payment_terms")      # осталось по умолчанию
    assert get_setting(mem_db, "preorder_lead_days") == config.PREORDER_LEAD_DAYS


def test_staff_invoice_pdf_has_requisites_terms_and_dates(web, mem_db, monkeypatch):
    """В счёте: ИНН/ОГРН, строка про НДС, условия оплаты, дата и адрес доставки.

    Текст PDF с кириллицей на машине без шрифта DejaVu превращается в квадратики
    (в Docker шрифт ставится сам), поэтому проверяем HTML, который уходит в PDF."""
    from pypdf import PdfReader
    from xhtml2pdf import pisa
    html_seen = []
    original = pisa.CreatePDF
    monkeypatch.setattr(pisa, "CreatePDF", lambda src, *a, **kw: (html_seen.append(src), original(src, *a, **kw))[1])
    monkeypatch.setattr(config, "SHOP_INN", "7700000001")
    monkeypatch.setattr(config, "SHOP_OGRN", "1027700000001")
    monkeypatch.setattr(config, "SHOP_VAT_NOTE", "НДС не облагается")
    buyer = make_user(mem_db, "b@test.ru")
    order = make_order(mem_db, buyer, seed_item(mem_db, stock=50), packs=1)
    order.desired_date = datetime(2026, 11, 4)
    order.delivery_method, order.delivery_address = "delivery", "Москва, Цветочная 5"
    mem_db.commit()
    as_admin(web, mem_db)
    r = web.get(f"/admin/orders/{order.id}/invoice")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    html = html_seen[-1]
    for needle in ("ИНН 7700000001", "ОГРН 1027700000001", "НДС не облагается",
                   "Условия оплаты: Предоплата 100%", "Нужно к: 04.11.2026",
                   "Доставка: Москва, Цветочная 5"):
        assert needle in html, needle
    assert "7700000001" in PdfReader(BytesIO(r.content)).pages[0].extract_text()


def test_admin_orders_excel_has_new_columns(web, mem_db):
    from openpyxl import load_workbook
    buyer = make_user(mem_db, "b@test.ru")
    order = make_order(mem_db, buyer, seed_item(mem_db, stock=50), packs=1)
    order.desired_date, order.delivery_method = datetime(2026, 11, 4), "pickup"
    mem_db.commit()
    as_admin(web, mem_db)
    ws = load_workbook(BytesIO(web.get("/admin/orders/export").content)).active
    header = [c.value for c in ws[1]]
    assert header[-4:] == ["Нужно к", "Получение", "Адрес доставки", "Номер в 1С"]
    assert [c.value for c in ws[2]][-4:-1] == ["04.11.2026", "Самовывоз", None]


def test_customers_excel_includes_all_clients_and_marketing_consent(web, mem_db):
    from openpyxl import load_workbook
    from app.services import legal
    yes = make_user(mem_db, "yes@test.ru", company_name="Согласен")
    make_user(mem_db, "no@test.ru", company_name="Не согласен")
    legal.record_consent(mem_db, yes, "marketing", True, None, source="profile")
    make_order(mem_db, yes, seed_item(mem_db, stock=50), packs=2)
    make_order(mem_db, yes, seed_item(mem_db, name="Другая", stock=50), packs=1, status="Отменён")
    as_admin(web, mem_db)
    rows = list(load_workbook(BytesIO(web.get("/admin/customers/export").content)).active.values)
    header = rows[0]
    data = {r[0]: dict(zip(header, r)) for r in rows[1:]}
    assert "Согласие на рекламу" in header and set(data) == {"Согласен", "Не согласен"}
    assert data["Согласен"]["Согласие на рекламу"] == "Да" and data["Согласен"]["Заказов"] == 1
    assert data["Не согласен"]["Согласие на рекламу"] == "Нет" and data["Не согласен"]["Заказов"] == 0


def test_staff_records_marketing_revocation_on_clients_request(web, mem_db):
    from app.models import ConsentLog
    from app.services import legal
    client = make_user(mem_db, "b@test.ru")
    legal.record_consent(mem_db, client, "marketing", True, None, source="profile")
    as_manager(web, mem_db)
    page = web.get(f"/admin/clients/{client.id}")
    assert "Клиент попросил отписать" in page.text
    token = csrf_from(page.text)
    web.post(f"/admin/clients/{client.id}/marketing_off", data={"csrf_token": token})
    mem_db.expire_all()
    assert mem_db.get(User, client.id).marketing_consent is False
    last = mem_db.query(ConsentLog).order_by(ConsentLog.id.desc()).first()
    assert (last.kind, last.granted, last.source) == ("marketing", False, "admin")
    assert "Клиент попросил отписать" not in web.get(f"/admin/clients/{client.id}").text
    web.post(f"/admin/clients/{client.id}/marketing_off", data={"csrf_token": token})      # повтор безопасен
    assert mem_db.query(ConsentLog).count() == 2
